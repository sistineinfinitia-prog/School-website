import openpyxl
import json
import os

xlsx_path = r'C:\Users\Rudolph\Downloads\2025 Syllabus tracker for HL_sl students.xlsx'
wb = openpyxl.load_workbook(xlsx_path, data_only=True)

# Also let's map each topic to its DP Year from the Scope and Sequence & SoW:
# DP1 (Year 12): B2 Programming, B1 Computational Thinking, B3 OOP, B4 ADTs (HL), A3 Databases, IA Criterion A & B & C
# DP2 (Year 13): A1 Computer Fundamentals, A4 Machine Learning, A2 Networks, IA Criterion D & E, Paper 3 / Case Study, Mocks & Revision

COHORT_MAPPING = {
    'A1 Computer fundamentals': 'DP2 (Year 2)',
    'A2 Networks': 'DP2 (Year 2)',
    'A3 Databases': 'DP1 (Year 1)',
    'A4 Machine learning': 'DP2 (Year 2)',
    'B1 Computational thinking': 'DP1 (Year 1)',
    'B2 Programming': 'DP1 (Year 1)',
    'B3 Object-oriented programming': 'DP1 (Year 1)',
    'B4 Abstract data types (HL only': 'DP1 (Year 1)'
}

HOURS_MAPPING = {
    'A1 Computer fundamentals': {'sl': 11, 'hl': 18, 'theme': 'Theme A: Concepts of Computer Science', 'paper': 'Paper 1'},
    'A2 Networks': {'sl': 11, 'hl': 18, 'theme': 'Theme A: Concepts of Computer Science', 'paper': 'Paper 1'},
    'A3 Databases': {'sl': 11, 'hl': 18, 'theme': 'Theme A: Concepts of Computer Science', 'paper': 'Paper 1'},
    'A4 Machine learning': {'sl': 5, 'hl': 18, 'theme': 'Theme A: Concepts of Computer Science', 'paper': 'Paper 1'},
    'B1 Computational thinking': {'sl': 5, 'hl': 5, 'theme': 'Theme B: Computational Thinking & Problem-Solving', 'paper': 'Paper 2'},
    'B2 Programming': {'sl': 40, 'hl': 42, 'theme': 'Theme B: Computational Thinking & Problem-Solving', 'paper': 'Paper 2'},
    'B3 Object-oriented programming': {'sl': 7, 'hl': 23, 'theme': 'Theme B: Computational Thinking & Problem-Solving', 'paper': 'Paper 2'},
    'B4 Abstract data types (HL only': {'sl': 0, 'hl': 23, 'theme': 'Theme B: Computational Thinking & Problem-Solving', 'paper': 'Paper 2'}
}

clean_units = []

for sheet_name in wb.sheetnames:
    if sheet_name in ['Intro', 'Progress']:
        continue
    sheet = wb[sheet_name]
    unit_info = HOURS_MAPPING.get(sheet_name, {})
    cohort_info = COHORT_MAPPING.get(sheet_name, 'DP1 & DP2')
    
    clean_title = sheet_name.replace(' (HL only', '').strip()
    code_prefix = clean_title.split()[0]
    title_rest = clean_title[len(code_prefix):].strip()
    
    curr_subtopic = None
    subtopics_dict = {}
    
    for r in range(4, sheet.max_row + 1):
        st = sheet.cell(r, 1).value
        if st and str(st).strip():
            curr_subtopic = str(st).strip()
            
        stmt = sheet.cell(r, 2).value
        lvl = sheet.cell(r, 3).value
        
        if stmt and str(stmt).strip():
            stmt_clean = str(stmt).strip()
            # If code is inside text e.g. "A1.1.1 Describe..."
            parts = stmt_clean.split(' ', 1)
            first_word = parts[0].strip()
            if any(first_word.startswith(p) for p in ['A', 'B']) and '.' in first_word:
                item_code = first_word
                item_text = parts[1].strip() if len(parts) > 1 else ''
            else:
                item_code = ''
                item_text = stmt_clean
                
            level_str = str(lvl).strip().upper() if lvl else ('HL' if 'HL' in sheet_name else 'SL')
            if 'HL ONLY' in stmt_clean.upper() or 'HL ONLY' in (curr_subtopic or '').upper():
                level_str = 'HL'
                
            if curr_subtopic not in subtopics_dict:
                subtopics_dict[curr_subtopic] = []
                
            subtopics_dict[curr_subtopic].append({
                'code': item_code,
                'statement': item_text,
                'full': stmt_clean,
                'level': level_str
            })

    subtopic_list = []
    for st_name, st_items in subtopics_dict.items():
        st_clean = st_name or clean_title
        st_parts = st_clean.split(' ', 1)
        st_code = st_parts[0] if ('.' in st_parts[0]) else ''
        st_title = st_parts[1] if ('.' in st_parts[0] and len(st_parts) > 1) else st_clean
        is_hl_subtopic = 'HL ONLY' in st_clean.upper() or all(it['level'] == 'HL' for it in st_items)
        
        subtopic_list.append({
            'code': st_code,
            'title': st_title,
            'fullName': st_clean,
            'isHLOnly': is_hl_subtopic,
            'statements': st_items
        })
        
    clean_units.append({
        'id': code_prefix.lower(),
        'code': code_prefix,
        'name': title_rest,
        'fullName': clean_title,
        'theme': unit_info.get('theme', ''),
        'paper': unit_info.get('paper', ''),
        'hoursSL': unit_info.get('sl', 0),
        'hoursHL': unit_info.get('hl', 0),
        'cohort': cohort_info,
        'isHLOnly': (unit_info.get('sl', 0) == 0),
        'subtopics': subtopic_list,
        'totalStatements': sum(len(st['statements']) for st in subtopic_list),
        'slStatements': sum(sum(1 for it in st['statements'] if it['level'] == 'SL') for st in subtopic_list),
        'hlStatements': sum(sum(1 for it in st['statements'] if it['level'] == 'HL') for st in subtopic_list)
    })

print(f"Parsed {len(clean_units)} syllabus units successfully:")
for u in clean_units:
    print(f"[{u['code']}] {u['name']} - {u['cohort']} | Paper: {u['paper']} | SL: {u['hoursSL']}h, HL: {u['hoursHL']}h | Statements: {u['totalStatements']} (SL:{u['slStatements']}, HL:{u['hlStatements']})")
    for st in u['subtopics']:
        print(f"    * {st['code']} {st['title']} ({len(st['statements'])} items)")

output_path = r'C:\Users\Rudolph\Documents\website\github repo\data\syllabus.json'
with open(output_path, 'w', encoding='utf-8') as f:
    json.dump(clean_units, f, indent=2, ensure_ascii=False)
print("Saved to", output_path)
