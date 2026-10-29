import openpyxl
import json
import os

xlsx_path = r'C:\Users\Rudolph\Downloads\2025 Syllabus tracker for HL_sl students.xlsx'
wb = openpyxl.load_workbook(xlsx_path, data_only=True)

UNIT_METADATA = {
    'A1': {
        'name': 'Computer Fundamentals',
        'theme': 'Theme A: Concepts of Computer Science',
        'paper': 'Paper 1',
        'cohort': 'DP2 (Year 2)',
        'hoursSL': 11,
        'hoursHL': 18,
        'hub': 'theme-a.html#a1',
        'description': 'Hardware architecture, Von Neumann CPU cycle, GPU functions, primary and secondary storage, logic gates, truth tables, operating systems, scheduling, interrupt handling, and compilation/interpretation translation.'
    },
    'A2': {
        'name': 'Networks',
        'theme': 'Theme A: Concepts of Computer Science',
        'paper': 'Paper 1',
        'cohort': 'DP2 (Year 2)',
        'hoursSL': 11,
        'hoursHL': 18,
        'hub': 'theme-a.html#a2',
        'description': 'Network fundamentals, digital infrastructures, devices and protocols, TCP/IP architecture, topologies, client-server models, segmentation, packet switching, IP addressing, static/dynamic routing, and network security countermeasures.'
    },
    'A3': {
        'name': 'Databases',
        'theme': 'Theme A: Concepts of Computer Science',
        'paper': 'Paper 1',
        'cohort': 'DP1 (Year 1)',
        'hoursSL': 11,
        'hoursHL': 18,
        'hub': 'theme-a.html#a3',
        'description': 'Relational database systems, schema modeling, Entity-Relationship Diagrams (ERDs), normalization (1NF to 3NF), denormalization, SQL queries and data updates, aggregate functions, transactions, and distributed data warehouses.'
    },
    'A4': {
        'name': 'Machine Learning',
        'theme': 'Theme A: Concepts of Computer Science',
        'paper': 'Paper 1',
        'cohort': 'DP2 (Year 2)',
        'hoursSL': 5,
        'hoursHL': 18,
        'hub': 'theme-a.html#a4',
        'description': 'Supervised, unsupervised, and reinforcement learning paradigms, hardware accelerators, data cleaning, dimensionality reduction, regression, classification, clustering, ANNs, CNNs, genetic algorithms, and societal ethics.'
    },
    'B1': {
        'name': 'Computational Thinking',
        'theme': 'Theme B: Computational Thinking & Problem-Solving',
        'paper': 'Paper 2',
        'cohort': 'DP1 (Year 1)',
        'hoursSL': 5,
        'hoursHL': 5,
        'hub': 'theme-b.html#b1',
        'description': 'Problem decomposition, abstraction, pattern recognition, algorithmic specification, flowchart construction and tracing, and systematic computational problem-solving.'
    },
    'B2': {
        'name': 'Programming',
        'theme': 'Theme B: Computational Thinking & Problem-Solving',
        'paper': 'Paper 2',
        'cohort': 'DP1 (Year 1)',
        'hoursSL': 40,
        'hoursHL': 42,
        'hub': 'theme-b.html#b2',
        'description': 'Core programming fundamentals in Python: variable types, substring operations, exception handling, arrays, dynamic lists, stacks, queues, sequence, selection, loops, functions, Big O analysis, linear/binary search, bubble/selection sorts, recursive algorithms, and file I/O.'
    },
    'B3': {
        'name': 'Object-Oriented Programming',
        'theme': 'Theme B: Computational Thinking & Problem-Solving',
        'paper': 'Paper 2',
        'cohort': 'DP1 (Year 1)',
        'hoursSL': 7,
        'hoursHL': 23,
        'hub': 'theme-b.html#b3',
        'description': 'OOP single class architecture: class definition, instantiation, static vs instance scope, encapsulation, and information hiding; extended in HL to multiple classes: inheritance, polymorphism, method overriding, abstraction, aggregation, and design patterns.'
    },
    'B4': {
        'name': 'Abstract Data Types (ADTs)',
        'theme': 'Theme B: Computational Thinking & Problem-Solving',
        'paper': 'Paper 2',
        'cohort': 'DP1 (Year 1)',
        'hoursSL': 0,
        'hoursHL': 23,
        'hub': 'theme-b.html#b4',
        'description': 'Higher Level only: core principles and properties of ADTs, singly, doubly, and circular linked lists, binary search trees (BST properties and operations), and sets.'
    }
}

syllabus_units = []

for sheet_name in wb.sheetnames:
    if sheet_name in ['Intro', 'Progress']:
        continue
    sheet = wb[sheet_name]
    
    code = sheet_name.split()[0].upper()
    meta = UNIT_METADATA.get(code, {})
    
    subtopics = {}
    curr_subtopic = None
    
    for r in range(4, sheet.max_row + 1):
        st = sheet.cell(r, 1).value
        if st and str(st).strip():
            curr_subtopic = str(st).strip()
            
        stmt = sheet.cell(r, 2).value
        lvl = sheet.cell(r, 3).value
        
        if stmt and str(stmt).strip():
            stmt_clean = str(stmt).strip()
            parts = stmt_clean.split(' ', 1)
            first_word = parts[0].strip()
            if any(first_word.startswith(p) for p in ['A', 'B']) and '.' in first_word:
                item_code = first_word
                item_text = parts[1].strip() if len(parts) > 1 else ''
            else:
                item_code = ''
                item_text = stmt_clean
                
            level_str = str(lvl).strip().upper() if lvl else ('HL' if 'HL' in sheet_name else 'SL')
            if 'HL ONLY' in stmt_clean.upper() or 'HL ONLY' in (curr_subtopic or '').upper():
                level_str = 'HL'
                
            if curr_subtopic not in subtopics:
                subtopics[curr_subtopic] = []
                
            subtopics[curr_subtopic].append({
                'code': item_code,
                'statement': item_text,
                'full': stmt_clean,
                'level': level_str
            })

    subtopic_list = []
    for st_name, st_items in subtopics.items():
        st_clean = st_name or meta.get('name', '')
        st_parts = st_clean.split(' ', 1)
        st_code = st_parts[0] if ('.' in st_parts[0]) else ''
        st_title = st_parts[1] if ('.' in st_parts[0] and len(st_parts) > 1) else st_clean
        is_hl_sub = 'HL ONLY' in st_clean.upper() or all(it['level'] == 'HL' for it in st_items)
        
        subtopic_list.append({
            'code': st_code,
            'title': st_title,
            'fullName': st_clean,
            'isHLOnly': is_hl_sub,
            'statements': st_items
        })
        
    syllabus_units.append({
        'id': code.lower(),
        'code': code,
        'name': meta.get('name', ''),
        'theme': meta.get('theme', ''),
        'paper': meta.get('paper', ''),
        'cohort': meta.get('cohort', ''),
        'hoursSL': meta.get('hoursSL', 0),
        'hoursHL': meta.get('hoursHL', 0),
        'isHLOnly': (meta.get('hoursSL', 0) == 0),
        'hub': meta.get('hub', ''),
        'description': meta.get('description', ''),
        'subtopics': subtopic_list,
        'totalStatements': sum(len(st['statements']) for st in subtopic_list),
        'slStatements': sum(sum(1 for it in st['statements'] if it['level'] == 'SL') for st in subtopic_list),
        'hlStatements': sum(sum(1 for it in st['statements'] if it['level'] == 'HL') for st in subtopic_list)
    })

# Write to data/syllabus.json
syllabus_out = r'C:\Users\Rudolph\Documents\website\github repo\data\syllabus.json'
with open(syllabus_out, 'w', encoding='utf-8') as f:
    json.dump(syllabus_units, f, indent=2, ensure_ascii=False)
print(f"Generated {syllabus_out} with {len(syllabus_units)} units.")

